from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.models.gallery_folder import GalleryFolder, GalleryType, UnicodeStatus
from app.models.toolkit_features import LibraryMatchPlan, PlanStatus
from app.services.library_match_service import LibraryMatchService


def gallery(path: Path, gallery_id: str, kind: GalleryType) -> GalleryFolder:
    return GalleryFolder(gallery_id, kind, path.name, path, UnicodeStatus.NFC, 0, 0, False, False, False, 1, "now")


def test_build_match_plan_respects_type_and_optional_cross_type(tmp_path: Path) -> None:
    a = gallery(tmp_path / "a" / "123 - A", "123", GalleryType.NORMAL)
    b = gallery(tmp_path / "b" / "Archive - 123 - B", "123", GalleryType.ARCHIVE)
    service = LibraryMatchService()
    assert service.build_plans([a], [b], tmp_path / "out") == []
    plans = service.build_plans([a], [b], tmp_path / "out", cross_type=True)
    assert len(plans) == 1
    assert plans[0].note == "跨类型 ID 匹配"
    assert plans[0].target_path == tmp_path / "out" / b.folder_name


def test_excel_requires_explicit_yes_and_uses_fixed_paths(tmp_path: Path) -> None:
    root_a = tmp_path / "a"; root_b = tmp_path / "b"; destination = tmp_path / "selected"
    source_a = root_a / "123 - A"; source_b = root_b / "123 - B"
    source_a.mkdir(parents=True); source_b.mkdir(parents=True)
    service = LibraryMatchService()
    plans = service.build_plans(
        [gallery(source_a, "123", GalleryType.NORMAL)],
        [gallery(source_b, "123", GalleryType.NORMAL)],
        destination,
    )
    workbook_path = tmp_path / "plan.xlsx"
    service.export_plan(plans, workbook_path)
    assert service.import_confirmed_plan(workbook_path, root_b, destination) == []

    workbook = load_workbook(workbook_path)
    workbook[service.SHEET_NAME]["A2"] = "是"
    workbook.save(workbook_path)
    imported = service.import_confirmed_plan(workbook_path, root_b, destination)
    assert len(imported) == 1 and imported[0].status is PlanStatus.READY
    service.execute(imported[0])
    assert not source_b.exists()
    assert (destination / "123 - B").is_dir()


def test_excel_rejects_source_or_target_outside_selected_roots(tmp_path: Path) -> None:
    root_a = tmp_path / "a"; root_b = tmp_path / "b"; destination = tmp_path / "out"
    source_a = root_a / "123 - A"; source_b = root_b / "123 - B"
    source_a.mkdir(parents=True); source_b.mkdir(parents=True)
    service = LibraryMatchService()
    plan = service.build_plans(
        [gallery(source_a, "123", GalleryType.NORMAL)],
        [gallery(source_b, "123", GalleryType.NORMAL)],
        destination,
    )
    workbook_path = tmp_path / "plan.xlsx"; service.export_plan(plan, workbook_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook[service.SHEET_NAME]
    sheet["A2"] = "是"
    sheet["H2"] = str(tmp_path / "outside" / "123 - B")
    sheet["I2"] = str(tmp_path / "also-outside" / "123 - B")
    workbook.save(workbook_path)
    imported = service.import_confirmed_plan(workbook_path, root_b, destination)
    assert imported[0].status is PlanStatus.CONFLICT
    assert "B 路径" in imported[0].note


def test_excel_duplicate_targets_are_conflicts(tmp_path: Path) -> None:
    root_a = tmp_path / "a"; root_b = tmp_path / "b"; destination = tmp_path / "out"
    a1 = root_a / "1 - A"; a2 = root_a / "2 - A"; b1 = root_b / "1 - B"; b2 = root_b / "2 - B"
    for path in (a1, a2, b1, b2): path.mkdir(parents=True)
    service = LibraryMatchService()
    plans = service.build_plans(
        [gallery(a1, "1", GalleryType.NORMAL), gallery(a2, "2", GalleryType.NORMAL)],
        [gallery(b1, "1", GalleryType.NORMAL), gallery(b2, "2", GalleryType.NORMAL)],
        destination,
    )
    workbook_path = tmp_path / "plan.xlsx"; service.export_plan(plans, workbook_path)
    workbook = load_workbook(workbook_path); sheet = workbook[service.SHEET_NAME]
    sheet["A2"] = "是"; sheet["A3"] = "是"; sheet["I3"] = sheet["I2"].value
    workbook.save(workbook_path)
    imported = service.import_confirmed_plan(workbook_path, root_b, destination)
    assert imported[0].status is PlanStatus.READY
    assert imported[1].status is PlanStatus.CONFLICT


def test_export_escapes_formula_like_gallery_names(tmp_path: Path) -> None:
    root_a = tmp_path / "a"
    root_b = tmp_path / "b"
    destination = tmp_path / "destination"
    source = root_b / "100 - normal"
    source.mkdir(parents=True)
    plan = LibraryMatchPlan(
        gallery_id="100",
        type_a="Normal",
        type_b="Normal",
        name_a="=HYPERLINK(\"https://example.invalid\")",
        name_b="+formula",
        path_a=root_a / "100 - normal",
        path_b=source,
        target_path=destination / source.name,
        note="@formula",
    )
    workbook_path = tmp_path / "plan.xlsx"

    LibraryMatchService().export_plan([plan], workbook_path)

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        row = next(
            workbook[LibraryMatchService.SHEET_NAME].iter_rows(
                min_row=2,
                max_row=2,
                values_only=True,
            )
        )
        assert row[4].startswith("'=")
        assert row[5].startswith("'+")
        assert row[10].startswith("'@")
        assert row[7] == str(source)
    finally:
        workbook.close()
