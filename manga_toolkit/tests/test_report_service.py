from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.models.gallery_status import GalleryReference, GalleryStatus, GalleryStatusRecord
from app.models.gallery_folder import GalleryFolder, GalleryType, UnicodeStatus
from app.models.library_analysis import DuplicateGroup, DuplicateKind, UnicodeAnalysisResult, UnicodeDuplicateGroup
from app.services.report_service import ReportService


def test_gallery_status_report_has_expected_sheets_and_no_token_column(tmp_path: Path) -> None:
    reference = GalleryReference(
        folder_path=tmp_path / "100 - title",
        source="metadata",
        gid=100,
        token="sensitive-token",
        gallery_url="https://e-hentai.org/g/100/sensitive-token/",
        local_title="title",
    )
    record = GalleryStatusRecord(
        reference=reference,
        status=GalleryStatus.UPDATE_AVAILABLE,
        checked_at=datetime.now(UTC),
        current_gid=200,
        current_token="new-sensitive-token",
        latest_url="https://e-hentai.org/g/200/new-sensitive-token/",
    )
    destination = tmp_path / "report.xlsx"

    ReportService().export_gallery_status([record], destination)
    workbook = load_workbook(destination, read_only=True)

    assert workbook.sheetnames == ["全部", "有新版本", "不可用画廊", "检查失败"]
    headers = [cell.value for cell in next(workbook["全部"].iter_rows())]
    assert "Token" not in "".join(str(value) for value in headers)


def test_duplicate_and_unicode_reports(tmp_path: Path) -> None:
    gallery = GalleryFolder(
        gallery_id="100",
        gallery_type=GalleryType.NORMAL,
        folder_name="100 - Cafe\u0301",
        path=tmp_path / "100 - Cafe\u0301",
        unicode_status=UnicodeStatus.NFD,
        file_count=1,
        folder_size=10,
        has_metadata=True,
        has_ametadata=False,
        has_comic_info=True,
        modified_ns=1,
        scanned_at=datetime.now(UTC).isoformat(),
    )
    duplicate = DuplicateGroup("100", DuplicateKind.NORMAL, (gallery, gallery))
    unicode_result = UnicodeAnalysisResult(
        [gallery],
        [UnicodeDuplicateGroup("100 - Café", (gallery,))],
    )
    duplicate_path = tmp_path / "duplicates.xlsx"
    unicode_path = tmp_path / "unicode.xlsx"

    ReportService().export_duplicates([duplicate], duplicate_path)
    ReportService().export_unicode(unicode_result, unicode_path)

    assert load_workbook(duplicate_path, read_only=True).sheetnames == ["重复ID"]
    assert load_workbook(unicode_path, read_only=True).sheetnames == ["编码检测", "Unicode重复"]


def test_export_generic_table(tmp_path: Path) -> None:
    destination = tmp_path / "table.xlsx"
    ReportService().export_table("目录检查", ["路径", "状态"], [[tmp_path, "正常"]], destination)
    assert load_workbook(destination, read_only=True).sheetnames == ["目录检查"]
