from __future__ import annotations

import shutil
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.models.gallery_folder import GalleryFolder
from app.models.toolkit_features import LibraryMatchPlan, PlanStatus


class LibraryMatchService:
    """Create and consume explicit Excel-backed A-ID -> B move plans."""

    SHEET_NAME = "匹配结果"
    HEADERS = (
        "执行",
        "ID",
        "A类型",
        "B类型",
        "A名称",
        "B名称",
        "A路径",
        "B路径",
        "目标路径",
        "状态",
        "备注",
    )
    AFFIRMATIVE = {"是", "yes", "y", "1", "true"}

    def build_plans(
        self,
        library_a: list[GalleryFolder],
        library_b: list[GalleryFolder],
        destination_root: Path,
        *,
        cross_type: bool = False,
    ) -> list[LibraryMatchPlan]:
        a_by_key: dict[tuple[str, str], list[GalleryFolder]] = defaultdict(list)
        for item in library_a:
            if item.gallery_id:
                key = ("*" if cross_type else item.gallery_type.value, item.gallery_id)
                a_by_key[key].append(item)
        plans: list[LibraryMatchPlan] = []
        for item_b in library_b:
            if not item_b.gallery_id:
                continue
            key = ("*" if cross_type else item_b.gallery_type.value, item_b.gallery_id)
            candidates = a_by_key.get(key)
            if not candidates:
                continue
            item_a = max(candidates, key=lambda item: len(item.folder_name))
            target = destination_root / item_b.folder_name
            plans.append(
                LibraryMatchPlan(
                    gallery_id=item_b.gallery_id,
                    type_a=item_a.gallery_type.value,
                    type_b=item_b.gallery_type.value,
                    name_a=item_a.folder_name,
                    name_b=item_b.folder_name,
                    path_a=item_a.path,
                    path_b=item_b.path,
                    target_path=target,
                    status=PlanStatus.CONFLICT if target.exists() else PlanStatus.READY,
                    note="跨类型 ID 匹配" if item_a.gallery_type is not item_b.gallery_type else "同类型 ID 匹配",
                )
            )
        return sorted(plans, key=lambda item: (int(item.gallery_id), item.type_b, item.name_b))

    def export_plan(self, plans: list[LibraryMatchPlan], destination: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.SHEET_NAME
        sheet.append(self.HEADERS)
        for plan in plans:
            sheet.append((
                "是" if plan.execute else "否",
                plan.gallery_id,
                plan.type_a,
                plan.type_b,
                self._excel_safe_text(plan.name_a),
                self._excel_safe_text(plan.name_b),
                str(plan.path_a),
                str(plan.path_b),
                str(plan.target_path),
                plan.status.value,
                self._excel_safe_text(plan.note),
            ))
        validation = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
        sheet.add_data_validation(validation)
        if sheet.max_row >= 2:
            validation.add(f"A2:A{sheet.max_row}")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = (10, 14, 12, 12, 36, 36, 60, 60, 60, 12, 24)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[chr(64 + index)].width = width
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)

    @staticmethod
    def _excel_safe_text(value: str) -> str:
        """Prevent gallery-controlled display text from becoming an Excel formula."""
        return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value

    def import_confirmed_plan(
        self,
        workbook_path: Path,
        library_b_root: Path,
        destination_root: Path,
    ) -> list[LibraryMatchPlan]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if self.SHEET_NAME not in workbook.sheetnames:
                raise ValueError(f"Excel 缺少工作表：{self.SHEET_NAME}")
            sheet = workbook[self.SHEET_NAME]
            headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            positions = {name: index for index, name in enumerate(headers)}
            missing = [name for name in self.HEADERS if name not in positions]
            if missing:
                raise ValueError("Excel 缺少列：" + "、".join(missing))
            plans: list[LibraryMatchPlan] = []
            seen_sources: set[str] = set()
            seen_targets: set[str] = set()
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                values = list(row)
                execute = str(values[positions["执行"]] or "").strip().casefold() in self.AFFIRMATIVE
                if not execute:
                    continue
                source = Path(str(values[positions["B路径"]] or "").strip())
                target = Path(str(values[positions["目标路径"]] or "").strip())
                status, note = self._validate_paths(source, target, library_b_root, destination_root)
                source_key = str(source).casefold()
                target_key = str(target).casefold()
                if source_key in seen_sources or target_key in seen_targets:
                    status = PlanStatus.CONFLICT
                    note = f"Excel 第 {row_number} 行与其他执行行重复源路径或目标路径"
                seen_sources.add(source_key)
                seen_targets.add(target_key)
                plans.append(
                    LibraryMatchPlan(
                        gallery_id=str(values[positions["ID"]] or ""),
                        type_a=str(values[positions["A类型"]] or ""),
                        type_b=str(values[positions["B类型"]] or ""),
                        name_a=str(values[positions["A名称"]] or ""),
                        name_b=str(values[positions["B名称"]] or ""),
                        path_a=Path(str(values[positions["A路径"]] or "")),
                        path_b=source,
                        target_path=target,
                        execute=True,
                        status=status,
                        note=note or str(values[positions["备注"]] or ""),
                    )
                )
            return plans
        finally:
            workbook.close()

    def execute(self, plan: LibraryMatchPlan) -> None:
        if not plan.execute or plan.status is not PlanStatus.READY:
            raise ValueError("该 Excel 记录未确认执行或存在冲突")
        if not plan.path_b.exists():
            raise FileNotFoundError(f"B 源路径不存在：{plan.path_b}")
        if plan.target_path.exists():
            raise FileExistsError(f"目标路径已存在：{plan.target_path}")
        plan.target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(plan.path_b), str(plan.target_path))

    @staticmethod
    def _validate_paths(
        source: Path,
        target: Path,
        library_b_root: Path,
        destination_root: Path,
    ) -> tuple[PlanStatus, str]:
        if not str(source) or not str(target):
            return PlanStatus.CONFLICT, "B 路径或目标路径为空"
        try:
            source_parent = source.resolve(strict=False).parent
            allowed_source = library_b_root.resolve(strict=False)
            target_parent = target.resolve(strict=False).parent
            allowed_target = destination_root.resolve(strict=False)
        except OSError as error:
            return PlanStatus.CONFLICT, f"路径解析失败：{error}"
        if source_parent != allowed_source:
            return PlanStatus.CONFLICT, "B 路径不是所选库 B 的一级项目"
        if target_parent != allowed_target:
            return PlanStatus.CONFLICT, "目标路径不是所选输出目录的一级项目"
        if not source.exists():
            return PlanStatus.SKIPPED, "B 源路径不存在"
        if target.exists():
            return PlanStatus.CONFLICT, "目标路径已存在"
        return PlanStatus.READY, ""
