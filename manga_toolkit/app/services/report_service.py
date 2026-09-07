from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.gallery_status import (
    ERROR_STATUSES,
    UNAVAILABLE_STATUSES,
    GalleryStatus,
    GalleryStatusRecord,
)
from app.models.library_analysis import DuplicateGroup, UnicodeAnalysisResult


class ReportService:
    def export_table(
        self,
        sheet_name: str,
        headers: list[str],
        rows: list[list[object]],
        destination: Path,
    ) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_sheet(
            workbook,
            sheet_name[:31] or "报告",
            headers,
            [[str(value) if isinstance(value, Path) else value for value in row] for row in rows],
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)

    def export_table_csv(
        self,
        headers: list[str],
        rows: list[list[object]],
        destination: Path,
    ) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        with destination.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows([
                [self._spreadsheet_safe(self._serializable(value)) for value in row]
                for row in rows
            ])

    def export_table_json(
        self,
        headers: list[str],
        rows: list[list[object]],
        destination: Path,
    ) -> None:
        destination.parent.mkdir(parents=True, exist_ok=True)
        payload = [
            {header: self._serializable(value) for header, value in zip(headers, row, strict=False)}
            for row in rows
        ]
        destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    @staticmethod
    def _serializable(value: object) -> object:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return str(value)

    @staticmethod
    def _spreadsheet_safe(value: object) -> object:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return f"'{value}"
        return value

    def export_duplicates(self, groups: list[DuplicateGroup], destination: Path) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        rows = [
            [
                group.gallery_id,
                group.kind.label,
                len(group.members),
                group.length_role(member),
                len(member.folder_name),
                member.gallery_type.value,
                member.folder_name,
                str(member.path),
                member.folder_size,
                member.unicode_status.value,
            ]
            for group in groups
            for member in group.members
        ]
        self._write_sheet(
            workbook,
            "重复ID",
            ["ID", "重复类型", "组内数量", "长短", "名称长度", "目录类型", "文件夹名", "路径", "大小(B)", "Unicode"],
            rows,
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)

    def export_unicode(self, result: UnicodeAnalysisResult, destination: Path) -> None:
        import unicodedata

        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_sheet(
            workbook,
            "编码检测",
            ["Unicode", "ID", "类型", "文件夹名", "NFC标准化名称", "路径"],
            [
                [
                    item.unicode_status.value,
                    item.gallery_id,
                    item.gallery_type.value,
                    item.folder_name,
                    unicodedata.normalize("NFC", item.folder_name),
                    str(item.path),
                ]
                for item in result.galleries
            ],
        )
        self._write_sheet(
            workbook,
            "Unicode重复",
            ["标准化名称", "组内数量", "ID", "类型", "Unicode", "原始文件夹名", "路径"],
            [
                [
                    group.normalized_name,
                    len(group.members),
                    member.gallery_id,
                    member.gallery_type.value,
                    member.unicode_status.value,
                    member.folder_name,
                    str(member.path),
                ]
                for group in result.duplicate_groups
                for member in group.members
            ],
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)

    def export_gallery_status(
        self,
        records: list[GalleryStatusRecord],
        destination: Path,
    ) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        self._write_all(workbook, records)
        self._write_updates(
            workbook,
            [record for record in records if record.status is GalleryStatus.UPDATE_AVAILABLE],
        )
        self._write_unavailable(
            workbook,
            [record for record in records if record.status in UNAVAILABLE_STATUSES],
        )
        self._write_failures(
            workbook,
            [record for record in records if record.status in ERROR_STATUSES],
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)

    def _write_all(self, workbook: Workbook, records: list[GalleryStatusRecord]) -> None:
        headers = [
            "状态", "本地文件夹", "来源", "本地GID", "最新GID", "本地标题",
            "最新标题", "本地页数", "最新页数", "本地大小", "最新大小",
            "最新URL", "最后检查时间", "备注", "错误",
        ]
        rows = [self._all_row(record) for record in records]
        self._write_sheet(workbook, "全部", headers, rows)

    def _write_updates(self, workbook: Workbook, records: list[GalleryStatusRecord]) -> None:
        headers = [
            "本地文件夹", "本地GID", "最新GID", "本地标题", "最新标题",
            "本地日文标题", "最新日文标题",
            "本地页数", "最新页数", "页数变化", "本地大小", "最新大小",
            "大小变化", "本地发布时间", "最新发布时间", "本地分类", "最新分类",
            "本地标签", "最新标签", "最新URL",
        ]
        rows = [
            [
                str(record.reference.folder_path),
                record.local_gid,
                record.current_gid,
                record.reference.local_title,
                record.current_metadata.title if record.current_metadata else "",
                record.local_metadata.title_jpn if record.local_metadata else "",
                record.current_metadata.title_jpn if record.current_metadata else "",
                self._local_filecount(record),
                record.current_metadata.filecount if record.current_metadata else None,
                record.page_delta,
                self._local_filesize(record),
                record.current_metadata.filesize if record.current_metadata else None,
                record.size_delta,
                record.local_metadata.posted if record.local_metadata else None,
                record.current_metadata.posted if record.current_metadata else None,
                record.local_metadata.category if record.local_metadata else "",
                record.current_metadata.category if record.current_metadata else "",
                ", ".join(record.local_metadata.tags) if record.local_metadata else "",
                ", ".join(record.current_metadata.tags) if record.current_metadata else "",
                record.latest_url,
            ]
            for record in records
        ]
        self._write_sheet(workbook, "有新版本", headers, rows)

    def _write_unavailable(self, workbook: Workbook, records: list[GalleryStatusRecord]) -> None:
        headers = ["文件夹", "GID", "状态", "URL", "备注", "最后检查时间"]
        rows = [
            [
                str(record.reference.folder_path),
                record.local_gid,
                record.status.label,
                record.reference.gallery_url,
                record.note,
                record.checked_at.isoformat(),
            ]
            for record in records
        ]
        self._write_sheet(workbook, "不可用画廊", headers, rows)

    def _write_failures(self, workbook: Workbook, records: list[GalleryStatusRecord]) -> None:
        headers = ["文件夹", "来源", "GID", "状态", "错误", "最后检查时间"]
        rows = [
            [
                str(record.reference.folder_path),
                record.reference.source,
                record.local_gid,
                record.status.label,
                record.error or record.note,
                record.checked_at.isoformat(),
            ]
            for record in records
        ]
        self._write_sheet(workbook, "检查失败", headers, rows)

    def _all_row(self, record: GalleryStatusRecord) -> list[object | None]:
        return [
            record.status.label,
            str(record.reference.folder_path),
            record.reference.source,
            record.local_gid,
            record.current_gid,
            record.reference.local_title,
            record.current_metadata.title if record.current_metadata else "",
            self._local_filecount(record),
            record.current_metadata.filecount if record.current_metadata else None,
            self._local_filesize(record),
            record.current_metadata.filesize if record.current_metadata else None,
            record.latest_url,
            record.checked_at.isoformat(),
            record.note,
            record.error,
        ]

    @staticmethod
    def _local_filecount(record: GalleryStatusRecord) -> int | None:
        return record.local_metadata.filecount if record.local_metadata else record.reference.local_filecount

    @staticmethod
    def _local_filesize(record: GalleryStatusRecord) -> int | None:
        return record.local_metadata.filesize if record.local_metadata else record.reference.local_filesize

    @staticmethod
    def _write_sheet(
        workbook: Workbook,
        title: str,
        headers: list[str],
        rows: list[list[object | None]],
    ) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([ReportService._spreadsheet_safe(value) for value in row])
        fill = PatternFill("solid", fgColor="DCE6F1")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            width = min(60, max(12, len(header) * 2 + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
