#!/usr/bin/env python3
"""在 Windows 上扫描 JHenTai 下载目录，并导出指定收藏分类为 XLSX。"""

from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any


METADATA_NAMES = {"metadata", "ametadata", "metadata.json", "ametadata.json"}
FOLDER_GID_PATTERN = re.compile(r"^(?:Archive\s*-\s*)?(\d+)\s*-", re.IGNORECASE)


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as file:
        value = json.load(file)
    if not isinstance(value, dict):
        raise ValueError("结果 JSON 的顶层结构不是对象")
    return value


def positive_int(value: Any) -> int | None:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def gid_from_metadata(path: Path) -> int | None:
    try:
        with path.open("r", encoding="utf-8-sig") as file:
            data = json.load(file)
    except (OSError, UnicodeError, json.JSONDecodeError):
        return None

    if not isinstance(data, dict):
        return None

    candidates = [data.get("gid")]
    for key in ("gallery", "archive"):
        nested = data.get(key)
        if isinstance(nested, dict):
            candidates.append(nested.get("gid"))

    for candidate in candidates:
        gid = positive_int(candidate)
        if gid is not None:
            return gid
    return None


def gid_from_folder_name(name: str) -> int | None:
    match = FOLDER_GID_PATTERN.match(name)
    return positive_int(match.group(1)) if match else None


def scan_download_folders(download_root: Path) -> tuple[dict[int, list[Path]], int]:
    """优先从 metadata/ametadata 取 GID，失败时再从文件夹名取。"""
    folders_by_gid: dict[int, list[Path]] = defaultdict(list)
    scanned_count = 0

    for current_root, directory_names, file_names in os.walk(download_root):
        current = Path(current_root)
        if current == download_root:
            continue

        scanned_count += 1
        lower_file_names = {name.lower(): name for name in file_names}
        gid = None

        for metadata_name in METADATA_NAMES:
            actual_name = lower_file_names.get(metadata_name)
            if actual_name is None:
                continue
            gid = gid_from_metadata(current / actual_name)
            if gid is not None:
                break

        if gid is None:
            gid = gid_from_folder_name(current.name)

        if gid is not None:
            folders_by_gid[gid].append(current)
            # 已识别为漫画文件夹后，无需继续扫描其中可能存在的子目录。
            directory_names[:] = []

    for gid, paths in folders_by_gid.items():
        folders_by_gid[gid] = sorted(set(paths), key=lambda path: str(path).casefold())

    return dict(folders_by_gid), scanned_count


def newest_title(item: dict[str, Any]) -> str:
    chain = item.get("updateChain")
    if isinstance(chain, list) and chain:
        last = chain[-1]
        if isinstance(last, dict):
            return str(last.get("title") or "")
    return ""


def update_chain_text(item: dict[str, Any]) -> str:
    chain = item.get("updateChain")
    if not isinstance(chain, list):
        return ""
    gids = []
    for entry in chain:
        if isinstance(entry, dict):
            gid = positive_int(entry.get("gid"))
            if gid is not None:
                gids.append(str(gid))
    return " → ".join(gids)


def matching_rows(
    items: list[dict[str, Any]],
    folders_by_gid: dict[int, list[Path]],
    category: str,
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    row_number = 0

    for item in items:
        if not isinstance(item, dict):
            continue
        gid = positive_int(item.get("gid"))
        if gid is None:
            continue

        matches: list[Path | None] = folders_by_gid.get(gid) or [None]
        for folder in matches:
            row_number += 1
            common = [
                row_number,
                "找到" if folder else "未找到",
                gid,
                folder.name if folder else "",
                str(folder) if folder else "",
                str(item.get("title") or ""),
            ]

            if category == "new":
                rows.append(common + [
                    positive_int(item.get("targetGid")) or "",
                    newest_title(item),
                    update_chain_text(item),
                    str(item.get("originalUrl") or ""),
                    str(item.get("targetUrl") or ""),
                    str(item.get("favoriteOutcome") or ""),
                ])
            elif category == "copyright":
                rows.append(common + [
                    str(item.get("copyrightHolder") or ""),
                    str(item.get("statusMessage") or ""),
                    str(item.get("originalUrl") or item.get("targetUrl") or ""),
                ])
            else:
                rows.append(common + [
                    str(item.get("statusMessage") or ""),
                    str(item.get("originalUrl") or item.get("targetUrl") or ""),
                ])

    return rows


def write_workbook(
    output_path: Path,
    source_json: Path,
    download_root: Path,
    report: dict[str, Any],
    folders_by_gid: dict[int, list[Path]],
) -> dict[str, dict[str, int]]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise RuntimeError(
            "缺少 openpyxl。请先在命令提示符执行：py -m pip install openpyxl"
        ) from error

    workbook = Workbook()
    workbook.remove(workbook.active)

    sheet_specs = [
        (
            "有新版图库",
            "newVersionGalleries",
            "new",
            ["序号", "匹配状态", "原始GID", "文件夹名称", "完整路径", "本地标题",
             "最新GID", "最新版标题", "更新链", "原始网址", "最新网址", "收藏结果"],
            "2F75B5",
        ),
        (
            "版权下架",
            "copyrightUnavailable",
            "copyright",
            ["序号", "匹配状态", "原始GID", "文件夹名称", "完整路径", "本地标题",
             "版权方", "下架说明", "原始网址"],
            "C00000",
        ),
        (
            "删除或不可用",
            "removedUnavailable",
            "removed",
            ["序号", "匹配状态", "原始GID", "文件夹名称", "完整路径", "本地标题",
             "状态说明", "原始网址"],
            "595959",
        ),
    ]

    summary: dict[str, dict[str, int]] = {}
    thin_gray = Side(style="thin", color="D9E2F3")

    for sheet_name, json_key, category, headers, header_color in sheet_specs:
        raw_items = report.get(json_key)
        items = raw_items if isinstance(raw_items, list) else []
        rows = matching_rows(items, folders_by_gid, category)

        worksheet = workbook.create_sheet(sheet_name)
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        worksheet.append(headers)

        for row in rows:
            worksheet.append(row)

        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor=header_color)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="medium", color=header_color))

        worksheet.row_dimensions[1].height = 28
        worksheet.auto_filter.ref = worksheet.dimensions

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=thin_gray)

            status_cell = row[1]
            if status_cell.value == "未找到":
                status_cell.fill = PatternFill("solid", fgColor="FFF2CC")
                status_cell.font = Font(color="9C6500", bold=True)
            else:
                status_cell.fill = PatternFill("solid", fgColor="E2F0D9")
                status_cell.font = Font(color="375623")

        widths = {
            1: 8,
            2: 11,
            3: 13,
            4: 48,
            5: 72,
            6: 60,
        }
        for column_index in range(7, len(headers) + 1):
            widths[column_index] = 28

        for column_index, width in widths.items():
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

        # 将 URL 单元格设置为可点击链接。
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                    cell.hyperlink = cell.value
                    cell.style = "Hyperlink"

        matched = sum(1 for row in rows if row[1] == "找到")
        missing = sum(1 for row in rows if row[1] == "未找到")
        summary[sheet_name] = {
            "json_items": len(items),
            "xlsx_rows": len(rows),
            "matched_rows": matched,
            "missing_items": missing,
        }

    workbook.properties.title = "JHenTai 收藏分类文件夹清单"
    workbook.properties.subject = f"来源：{source_json.name}；下载目录：{download_root}"
    workbook.save(output_path)
    return summary


def main() -> int:
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox
    except ImportError:
        print("当前 Python 缺少 tkinter，建议安装 python.org 提供的 Windows Python。")
        return 1

    root = tk.Tk()
    root.withdraw()

    try:
        json_name = filedialog.askopenfilename(
            title="选择 jh-favorite-classified-cat-6.json",
            filetypes=[("JSON 文件", "*.json"), ("所有文件", "*.*")],
        )
        if not json_name:
            return 0

        download_name = filedialog.askdirectory(
            title="选择 JHenTai 下载根目录（包含漫画文件夹/分组文件夹）"
        )
        if not download_name:
            return 0

        json_path = Path(json_name)
        download_root = Path(download_name)
        default_output = json_path.with_name(
            f"{json_path.stem}-folders.xlsx"
        )

        output_name = filedialog.asksaveasfilename(
            title="保存分类 Excel",
            initialdir=str(default_output.parent),
            initialfile=default_output.name,
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not output_name:
            return 0

        print(f"正在扫描下载目录：{download_root}")
        folders_by_gid, scanned_count = scan_download_folders(download_root)
        print(
            f"扫描完成：检查 {scanned_count} 个目录，识别出 "
            f"{len(folders_by_gid)} 个不同 GID。"
        )

        report = load_json(json_path)
        summary = write_workbook(
            Path(output_name),
            json_path,
            download_root,
            report,
            folders_by_gid,
        )

        lines = [f"Excel 已生成：\n{output_name}\n"]
        for sheet_name, values in summary.items():
            lines.append(
                f"{sheet_name}：JSON {values['json_items']} 项，"
                f"找到 {values['matched_rows']} 行，"
                f"未找到 {values['missing_items']} 项"
            )

        result_text = "\n".join(lines)
        print(result_text)
        messagebox.showinfo("处理完成", result_text)
        return 0
    except Exception as error:
        message = f"处理失败：\n{error}"
        print(message, file=sys.stderr)
        messagebox.showerror("处理失败", message)
        return 1
    finally:
        root.destroy()


if __name__ == "__main__":
    raise SystemExit(main())
