#!/usr/bin/env python3
"""根据 JHenTai 分类 XLSX，在 Windows 上预演、复制或移动漫画文件夹。"""

from __future__ import annotations

import argparse
import csv
import os
import shutil
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


SHEET_DESTINATIONS = {
    "有新版图库": "01_有新版图库",
    "版权下架": "02_版权下架",
    "删除或不可用": "03_删除或不可用",
}


@dataclass(frozen=True)
class Assignment:
    category: str
    gid: str
    folder_name: str
    excel_path: str
    row_number: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="根据分类 Excel 将 JHenTai 文件夹收集到三个目录。默认仅预演。"
    )
    parser.add_argument("--xlsx", help="分类 Excel 路径；省略时弹窗选择")
    parser.add_argument("--source-root", help="JHenTai 下载根目录；省略时弹窗选择")
    parser.add_argument("--destination-root", help="分类输出根目录；省略时弹窗选择")
    parser.add_argument(
        "--mode",
        choices=("preview", "copy", "move"),
        default="preview",
        help="preview=只预演；copy=复制；move=移动。默认 preview",
    )
    return parser.parse_args()


def choose_paths(args: argparse.Namespace) -> tuple[Path, Path, Path]:
    if args.xlsx and args.source_root and args.destination_root:
        return Path(args.xlsx), Path(args.source_root), Path(args.destination_root)

    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError as error:
        raise RuntimeError(
            "缺少 tkinter，请通过命令行同时提供 --xlsx、--source-root、--destination-root。"
        ) from error

    root = tk.Tk()
    root.withdraw()
    try:
        xlsx = args.xlsx or filedialog.askopenfilename(
            title="选择 jh-favorite-classified-*-folders.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")],
        )
        if not xlsx:
            raise KeyboardInterrupt

        source_root = args.source_root or filedialog.askdirectory(
            title="选择 JHenTai 下载根目录（例如 T:\\JH）"
        )
        if not source_root:
            raise KeyboardInterrupt

        destination_root = args.destination_root or filedialog.askdirectory(
            title="选择分类输出根目录（不能位于 JHenTai 下载根目录内）"
        )
        if not destination_root:
            raise KeyboardInterrupt

        return Path(xlsx), Path(source_root), Path(destination_root)
    finally:
        root.destroy()


def normalized_path(path: Path) -> str:
    return os.path.normcase(os.path.abspath(str(path)))


def is_within(path: Path, root: Path) -> bool:
    try:
        return os.path.commonpath([normalized_path(path), normalized_path(root)]) == normalized_path(root)
    except ValueError:
        # Windows 不同盘符会触发 ValueError。
        return False


def validate_roots(xlsx: Path, source_root: Path, destination_root: Path) -> None:
    if not xlsx.is_file():
        raise FileNotFoundError(f"Excel 不存在：{xlsx}")
    if xlsx.suffix.lower() != ".xlsx":
        raise ValueError("输入文件必须是 .xlsx")
    if not source_root.is_dir():
        raise NotADirectoryError(f"下载根目录不存在：{source_root}")
    if not destination_root.is_dir():
        raise NotADirectoryError(f"分类输出根目录不存在：{destination_root}")
    if is_within(destination_root, source_root):
        raise ValueError("分类输出根目录不能等于或位于 JHenTai 下载根目录内部。")


def read_assignments(xlsx: Path) -> list[Assignment]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "缺少 openpyxl，请先执行：py -m pip install openpyxl"
        ) from error

    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    assignments: list[Assignment] = []

    try:
        for sheet_name in SHEET_DESTINATIONS:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Excel 缺少工作表：{sheet_name}")

            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_values = next(rows)
            except StopIteration as error:
                raise ValueError(f"工作表为空：{sheet_name}") from error

            headers = {
                str(value).strip(): index
                for index, value in enumerate(header_values)
                if value is not None
            }
            required = {"匹配状态", "原始GID", "文件夹名称", "完整路径"}
            missing = required - set(headers)
            if missing:
                raise ValueError(
                    f"工作表“{sheet_name}”缺少列：{', '.join(sorted(missing))}"
                )

            for row_number, values in enumerate(rows, start=2):
                status = str(values[headers["匹配状态"]] or "").strip()
                if status != "找到":
                    continue

                gid_value = values[headers["原始GID"]]
                gid = str(int(gid_value)) if isinstance(gid_value, (int, float)) else str(gid_value or "").strip()
                folder_name = str(values[headers["文件夹名称"]] or "").strip()
                excel_path = str(values[headers["完整路径"]] or "").strip()

                if not folder_name or not excel_path:
                    continue

                assignments.append(
                    Assignment(
                        category=sheet_name,
                        gid=gid,
                        folder_name=folder_name,
                        excel_path=excel_path,
                        row_number=row_number,
                    )
                )
    finally:
        workbook.close()

    return assignments


def build_folder_name_index(source_root: Path, wanted_names: set[str]) -> dict[str, list[Path]]:
    """仅用于 Excel 路径失效时的回退定位。"""
    index: dict[str, list[Path]] = {}
    wanted_casefold = {name.casefold(): name for name in wanted_names}

    for current_root, directory_names, _ in os.walk(source_root):
        current = Path(current_root)
        for directory_name in list(directory_names):
            key = directory_name.casefold()
            if key in wanted_casefold:
                original_name = wanted_casefold[key]
                index.setdefault(original_name, []).append(current / directory_name)
                # 漫画目录通常不需要继续递归。
                directory_names.remove(directory_name)

    return index


def resolve_sources(
    assignments: Iterable[Assignment], source_root: Path
) -> tuple[list[tuple[Assignment, Path | None, str]], dict[str, list[Path]]]:
    assignments = list(assignments)
    wanted_names = {item.folder_name for item in assignments}
    fallback_index: dict[str, list[Path]] | None = None
    resolved: list[tuple[Assignment, Path | None, str]] = []

    for assignment in assignments:
        excel_source = Path(assignment.excel_path)

        if excel_source.is_dir() and is_within(excel_source, source_root):
            resolved.append((assignment, excel_source, "excel_path"))
            continue

        direct_source = source_root / assignment.folder_name
        if direct_source.is_dir():
            resolved.append((assignment, direct_source, "source_root_name"))
            continue

        if fallback_index is None:
            print("部分 Excel 路径失效，正在下载根目录中按文件夹名称回退查找……")
            fallback_index = build_folder_name_index(source_root, wanted_names)

        matches = fallback_index.get(assignment.folder_name, [])
        if len(matches) == 1:
            resolved.append((assignment, matches[0], "recursive_name_search"))
        elif len(matches) > 1:
            resolved.append((assignment, None, "ambiguous_multiple_matches"))
        else:
            resolved.append((assignment, None, "not_found"))

    return resolved, fallback_index or {}


def deduplicate_and_find_conflicts(
    resolved: Iterable[tuple[Assignment, Path | None, str]]
) -> tuple[list[tuple[Assignment, Path | None, str]], set[str]]:
    unique: dict[tuple[str, str], tuple[Assignment, Path | None, str]] = {}
    categories_by_source: dict[str, set[str]] = {}

    for item in resolved:
        assignment, source, method = item
        source_key = normalized_path(source) if source else f"missing:{assignment.category}:{assignment.gid}:{assignment.folder_name}"
        unique[(source_key, assignment.category)] = item
        if source:
            categories_by_source.setdefault(source_key, set()).add(assignment.category)

    conflicts = {
        source_key
        for source_key, categories in categories_by_source.items()
        if len(categories) > 1
    }
    return list(unique.values()), conflicts


def safe_destination(destination_root: Path, category: str, source: Path) -> Path:
    category_folder = destination_root / SHEET_DESTINATIONS[category]
    return category_folder / source.name


def confirm_real_operation(mode: str, count: int, destination_root: Path) -> None:
    if mode == "preview":
        return

    expected = mode.upper()
    print()
    print(f"即将以 {mode} 模式处理最多 {count} 个文件夹。")
    print(f"目标根目录：{destination_root}")
    print("不会覆盖或合并已存在的同名目标。")
    answer = input(f"确认后请输入 {expected}：").strip()
    if answer != expected:
        raise KeyboardInterrupt


def execute(
    items: list[tuple[Assignment, Path | None, str]],
    conflicts: set[str],
    source_root: Path,
    destination_root: Path,
    mode: str,
) -> list[dict[str, Any]]:
    logs: list[dict[str, Any]] = []

    if mode != "preview":
        for folder_name in SHEET_DESTINATIONS.values():
            (destination_root / folder_name).mkdir(parents=True, exist_ok=True)

    for index, (assignment, source, resolve_method) in enumerate(items, start=1):
        destination = safe_destination(destination_root, assignment.category, source) if source else None
        status = ""
        detail = ""

        try:
            if source is None:
                status = "skipped_not_found"
                detail = resolve_method
            elif normalized_path(source) in conflicts:
                status = "skipped_category_conflict"
                detail = "同一源文件夹出现在多个分类中"
            elif not source.is_dir():
                status = "skipped_not_found"
                detail = "执行时源文件夹已不存在"
            elif source.is_symlink():
                status = "skipped_symlink"
                detail = "为避免误操作，不处理符号链接目录"
            elif not is_within(source, source_root):
                status = "skipped_outside_source_root"
                detail = "源路径不在所选下载根目录内"
            elif destination is None:
                status = "error"
                detail = "无法计算目标路径"
            elif destination.exists():
                status = "skipped_destination_exists"
                detail = "目标已存在；未覆盖、未合并"
            elif mode == "preview":
                status = "planned"
                detail = f"将通过 {resolve_method} 定位"
            elif mode == "copy":
                shutil.copytree(source, destination)
                status = "copied"
            elif mode == "move":
                shutil.move(str(source), str(destination))
                status = "moved"
            else:
                raise ValueError(f"未知模式：{mode}")
        except Exception as error:  # 单项失败不终止整个批次
            status = "error"
            detail = f"{type(error).__name__}: {error}"

        logs.append(
            {
                "index": index,
                "mode": mode,
                "category": assignment.category,
                "gid": assignment.gid,
                "folder_name": assignment.folder_name,
                "excel_row": assignment.row_number,
                "source": str(source or assignment.excel_path),
                "destination": str(destination or ""),
                "resolve_method": resolve_method,
                "status": status,
                "detail": detail,
            }
        )
        print(f"[{index}/{len(items)}] {status}: {assignment.gid} {assignment.folder_name}")

    return logs


def write_log(logs: list[dict[str, Any]], xlsx: Path, mode: str) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = xlsx.with_name(f"jh_folder_{mode}_log_{timestamp}.csv")
    fieldnames = [
        "index", "mode", "category", "gid", "folder_name", "excel_row",
        "source", "destination", "resolve_method", "status", "detail",
    ]
    with log_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(logs)
    return log_path


def summarize(logs: list[dict[str, Any]]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for item in logs:
        status = str(item["status"])
        counts[status] = counts.get(status, 0) + 1
    return counts


def main() -> int:
    args = parse_args()
    try:
        xlsx, source_root, destination_root = choose_paths(args)
        xlsx = xlsx.resolve()
        source_root = source_root.resolve()
        destination_root = destination_root.resolve()

        validate_roots(xlsx, source_root, destination_root)
        assignments = read_assignments(xlsx)
        if not assignments:
            raise ValueError("Excel 中没有“匹配状态=找到”的可处理记录。")

        resolved, _ = resolve_sources(assignments, source_root)
        items, conflicts = deduplicate_and_find_conflicts(resolved)

        print(f"读取记录：{len(assignments)}")
        print(f"去重后文件夹任务：{len(items)}")
        print(f"跨分类冲突：{len(conflicts)}")
        print(f"模式：{args.mode}")

        confirm_real_operation(args.mode, len(items), destination_root)
        logs = execute(items, conflicts, source_root, destination_root, args.mode)
        log_path = write_log(logs, xlsx, args.mode)
        counts = summarize(logs)

        print("\n处理完成：")
        for status, count in sorted(counts.items()):
            print(f"  {status}: {count}")
        print(f"日志：{log_path}")

        if args.mode == "preview":
            print("\n确认日志无误后，重新运行并添加 --mode move。")
        return 0
    except KeyboardInterrupt:
        print("已取消，未继续执行。")
        return 130
    except Exception as error:
        print(f"处理失败：{type(error).__name__}: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
